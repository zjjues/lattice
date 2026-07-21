from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/home/cring/lattice/entity-generator")
OUT = ROOT / "external/fake_lattice_系统说明.docx"
IMG_DIR = Path("/tmp/fake_lattice_doc")


def set_run_font(run, *, size: int | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para(paragraph, *, alignment=None, space_before=0, space_after=0, line_spacing=1.2):
    if alignment is not None:
        paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading("", level=level)
    p.clear()
    r = p.add_run(text)
    size = {1: 20, 2: 15, 3: 13}.get(level, 12)
    set_run_font(r, size=size, bold=True)
    set_para(p, space_before=4, space_after=4, line_spacing=1.0)
    return p


def add_body(doc: Document, text: str, *, size: int = 11, bold: bool = False, color: str | None = None):
    p = doc.add_paragraph()
    set_para(p, space_before=0, space_after=4, line_spacing=1.18)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_bullet(doc: Document, text: str, level: int = 0, size: int = 11):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.18 * level)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_numbered(doc: Document, text: str, size: int = 11):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def add_table(doc: Document, rows: list[list[str]], widths: list[float] | None = None, header_fill=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for col_idx, width in enumerate(widths):
            for row in table.rows:
                row.cells[col_idx].width = Inches(width)
    for r_idx, row_data in enumerate(rows):
        for c_idx, text in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            set_para(p, space_before=0, space_after=0, line_spacing=1.05)
            r = p.add_run(text)
            set_run_font(r, size=11, bold=(r_idx == 0))
            if r_idx == 0:
                for run in p.runs:
                    set_run_font(run, size=11, bold=True)
    return table


def add_screenshot(doc: Document, image_name: str, caption: str, note_lines: list[str] | None = None):
    img = IMG_DIR / image_name
    if not img.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, space_before=4, space_after=2, line_spacing=1.0)
    p.add_run().add_picture(str(img), width=Inches(6.9))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(cap, space_before=2, space_after=2, line_spacing=1.0)
    r = cap.add_run(caption)
    set_run_font(r, size=10, bold=True, color="666666")
    if note_lines:
        for line in note_lines:
            add_body(doc, line, size=10)


def new_page(doc: Document):
    doc.add_page_break()


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(11)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, space_before=120, space_after=20, line_spacing=1.0)
    r = p.add_run("Fake Lattice 本地指控系统说明")
    set_run_font(r, size=24, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, space_before=8, space_after=8, line_spacing=1.0)
    r = p.add_run("需求分析 / 定位 / 功能说明 / 当前进展 / 开发计划 / 典型应用场景")
    set_run_font(r, size=12, color="555555")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, space_before=10, space_after=2, line_spacing=1.0)
    r = p.add_run("基于当前代码工程整理")
    set_run_font(r, size=11, color="777777")
    new_page(doc)

    # 1. 需求分析
    add_heading(doc, "1. 需求分析", 1)
    add_body(doc, "这套 fake lattice 的出现，不是为了复刻 Anduril 官方 Lattice，而是为了解决“真实平台不可得、但又需要把 Lattice 风格闭环讲清楚”这个问题。", size=11)
    add_bullet(doc, "真实 Lattice 不开源，部分地区也无法获得 endpoint、token、client id / secret、真实 UI。", size=11)
    add_bullet(doc, "演示需要可重复、可本地运行、可重置的数据闭环，不能依赖外部网络或现场授权。", size=11)
    add_bullet(doc, "如果只展示静态页面，无法说明“数据接入 -> 告警 -> 建议 -> 任务 -> 回流”的完整链路。", size=11)
    add_bullet(doc, "因此需要一套本地可跑的教学版系统，既能说明平台思想，也能说明工程实现。", size=11)

    add_heading(doc, "2. 项目定位", 1)
    add_body(doc, "fake lattice 的定位可以概括为三句话：", size=11, bold=True)
    add_bullet(doc, "它不是官方 Lattice，而是一个本地演示系统。", size=11)
    add_bullet(doc, "它面向的是教学、讲解、研发验证，不是生产部署。", size=11)
    add_bullet(doc, "它的核心价值不是功能堆叠，而是把 Lattice 风格的作战语义链路讲顺。", size=11)

    add_table(doc, [
        ["维度", "fake lattice 的定位", "明确不做什么"],
        ["平台属性", "本地可运行的演示平台", "不提供官方后端、真实权限、真实 endpoint"],
        ["交互对象", "目标、告警、建议、任务、对象", "不模拟完整实战系统和真实硬件协议"],
        ["智能部分", "本地知识库 + Dify 生成建议", "不声称是实战级智能调度引擎"],
        ["使用方式", "一键启动、回放轨迹、可重置", "不依赖外部服务才能演示"],
    ], widths=[1.05, 3.25, 2.8])

    new_page(doc)

    # 3. 功能说明
    add_heading(doc, "3. 功能说明", 1)
    add_body(doc, "当前代码工程里的 fake lattice，已经形成一套完整的本地闭环。下面按模块看它具体做什么。", size=11)

    add_table(doc, [
        ["模块", "当前能力", "对应代码"],
        ["接入层", "接收模拟事件，支持 POST 注入和轨迹回放", "mock_server.py / tools/send_sim_event.py / tools/replay_sim_events.py"],
        ["平台模型层", "统一维护 entities、alerts、recommendations、tasks、objects", "mock_server.py"],
        ["推理层", "本地规则告警、本地知识库检索、Dify Workflow 生成建议", "mock_server.py + kb/ + .env.local"],
        ["执行层", "sim_agent 模拟任务执行和状态回流", "tools/sim_agent.py"],
        ["前端层", "地图、摘要、告警、建议、任务、威胁目标联动展示", "static/index.html / static/app.js / static/style.css"],
        ["运维层", "重置、启动、停止、runbook、彩排清单、命令速查", "tools/reset_demo_state.py / scripts/*.sh / RUN_FAKE_LATTICE.md"],
    ], widths=[1.2, 3.3, 3.1])

    add_body(doc, "从工程角度看，这不是一个单页演示壳，而是一套可重复演示的本地指控原型。", size=11, bold=True)

    new_page(doc)

    # 4. 当前进展
    add_heading(doc, "4. 当前进展", 1)
    add_body(doc, "当前系统已经足够支撑演示和讲解，而且能够把“目标出现之后发生了什么”讲完整。", size=11)

    add_table(doc, [
        ["已经实现", "当前边界"],
        ["本地可运行后端：entities、alerts、recommendations、tasks、objects、WebSocket。", "不是 Anduril 官方 Lattice，没有真实后端、权限和认证体系。"],
        ["两类演示场景：警戒区入侵 / 盘旋滞留。", "当前地图是离线态势图，不是标准在线 GIS 瓦片地图。"],
        ["本地知识库参与 recommendation 生成，可接本地 Dify Workflow，也可本地回退。", "数据接入仍是模拟事件和轨迹回放，不是真实硬件。"],
        ["一键重置、启动、停止、runbook、彩排清单、命令速查。", "建议生成偏演示辅助，不是实战验证的智能调度系统。"],
    ], widths=[3.65, 3.45])

    add_heading(doc, "截图：空态界面", 2)
    add_screenshot(
        doc,
        "01-empty.png",
        "图 1  空态界面",
        [
            "启动后页面默认呈现空态：顶部摘要为 0，地图、告警、建议、任务均未加载任何目标。",
            "这张图适合说明系统的初始状态和整体布局。",
        ],
    )
    new_page(doc)

    add_heading(doc, "截图：告警态界面", 2)
    add_screenshot(
        doc,
        "02-alert.png",
        "图 2  告警态界面",
        [
            "注入模拟目标后，页面会出现目标点、轨迹和高风险告警。",
            "此时仍然没有建议和任务，说明系统先做态势识别，再进入处置决策。",
        ],
    )
    new_page(doc)

    add_heading(doc, "截图：建议与任务态界面", 2)
    add_screenshot(
        doc,
        "03-task.png",
        "图 3  建议与任务态界面",
        [
            "在告警基础上生成建议并创建任务后，系统会把 recommendation、task、状态回流串起来。",
            "这张图最适合解释 fake lattice 的闭环价值。",
        ],
    )

    new_page(doc)

    # 5. 典型应用场景
    add_heading(doc, "5. 典型应用场景", 1)
    add_body(doc, "当前 fake lattice 的用途不是泛化成一个大平台，而是围绕几个典型场景把链路讲清楚。", size=11)
    add_bullet(doc, "目标进入警戒区：展示告警如何触发，以及如何从告警进入处置建议。", size=11)
    add_bullet(doc, "目标盘旋滞留：展示风险不只来自“闯入”，也可能来自“持续观察”和“异常滞留”。", size=11)
    add_bullet(doc, "多源数据接入：展示不同来源的数据最终如何映射到统一 entity 模型。", size=11)
    add_bullet(doc, "知识库 + 大模型辅助研判：展示本地 SOP / 预案如何进入 recommendation 链路。", size=11)
    add_bullet(doc, "任务回流与执行状态跟踪：展示从“建议”到“任务”再到“回流”的闭环。", size=11)

    add_heading(doc, "6. 下一步开发计划", 1)
    add_numbered(doc, "继续完善模拟数据源，让更多传感器、更多轨迹形态、更多事件类型能够复用同一套接入层。", size=11)
    add_numbered(doc, "把 recommendation 的知识库引用做得更结构化，让建议来源、依据摘要和引用文档更容易追踪。", size=11)
    add_numbered(doc, "继续收敛前端布局，让当前事件、告警、建议、任务在更小的认知负担下呈现。", size=11)
    add_numbered(doc, "补充更多可讲解的演示场景，例如 AIS、传感器联动、多目标同时告警等。", size=11)
    add_numbered(doc, "如果后续要更贴近真实接入，再单独做协议适配层，而不是破坏当前本地演示闭环。", size=11)

    new_page(doc)

    # 7. 代码工程对应关系
    add_heading(doc, "7. 代码工程对应关系", 1)
    add_table(doc, [
        ["文件 / 目录", "作用"],
        ["mock_server.py", "核心本地 API，负责 entities、alerts、recommendations、tasks、objects、KB、Dify 适配。"],
        ["static/", "前端监控台，负责地图、摘要、当前事件、告警、建议、任务、威胁目标展示。"],
        ["tools/send_sim_event.py", "单次事件注入脚本，用于快速制造一个告警场景。"],
        ["tools/replay_sim_events.py", "轨迹回放脚本，用于制造连续变化的演示场景。"],
        ["tools/sim_agent.py", "模拟任务执行器，推进 task 状态并写回结果。"],
        ["tools/reset_demo_state.py", "一键清空演示状态。"],
        ["scripts/start_demo_env.sh / stop_demo_env.sh", "一键启动和停止演示环境。"],
        ["RUN_FAKE_LATTICE.md", "演示 runbook。"],
        ["docs/demo-final-checklist.md / docs/demo-cheatsheet.md", "演示彩排清单和现场速查。"],
        ["kb/dify_samples/", "可直接导入 Dify 的示例知识文件。"],
    ], widths=[2.6, 4.8])

    add_body(doc, "这份文档对应的是当前“已经能演示”的 fake lattice，而不是规划中的理想形态。后续如果代码继续变化，只需要更新这一节即可。", size=11)

    doc.save(str(OUT))


if __name__ == "__main__":
    main()
