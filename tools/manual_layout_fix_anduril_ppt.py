from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE


SRC_DEFAULT = Path("external/anduril-sdk介绍_原文仅排版修正版.pptx")
DST_DEFAULT = Path("external/anduril-sdk介绍_原文仅排版修正版.pptx")


def text_of(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def set_font(shape, size: float, bold: bool | None = None, name: str = "Microsoft YaHei"):
    if not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.name = name
            r.font.size = Pt(size)
            if bold is not None:
                r.font.bold = bold


def move(shape, left=None, top=None, width=None, height=None):
    if left is not None:
        shape.left = Inches(left)
    if top is not None:
        shape.top = Inches(top)
    if width is not None:
        shape.width = Inches(width)
    if height is not None:
        shape.height = Inches(height)


def find_by_text(slide, needle: str):
    for sh in slide.shapes:
        if text_of(sh).strip() == needle:
            return sh
    raise KeyError(f"shape not found: {needle!r}")


def find_contains(slide, needle: str):
    for sh in slide.shapes:
        if needle in text_of(sh):
            return sh
    raise KeyError(f"shape not found containing: {needle!r}")


def fix_slide_5(slide):
    # Rebuild the product area as a 3x2 grid of wider cards.
    cards = [
        ("Fury", "协同空中作战", 32, 33, 34, 35, 36, 0.60, 4.72),
        ("Roadrunner", "可回收防空节点", 37, 38, 39, 40, 41, 4.05, 4.72),
        ("Barracuda", "低成本规模化打击", 42, 43, 44, 45, 46, 7.50, 4.72),
        ("Ghost", "小型侦察无人机", 47, 48, 49, 50, 51, 0.60, 5.66),
        ("Dive-LD", "海上无人潜航器", 52, 53, 54, 55, 56, 4.05, 5.66),
        ("Anvil", "反无人机拦截", 57, 58, 59, 60, 61, 7.50, 5.66),
    ]

    # Make the explanatory ribbon a touch larger and lower so the cards can breathe.
    try:
        move(slide.shapes[68], left=0.62, top=6.52, width=11.40, height=0.44)
    except Exception:
        pass
    set_font(slide.shapes[68], 15.5, False)

    for title, desc, bg_idx, line_idx, img_idx, t_idx, d_idx, x, y in cards:
        bg = slide.shapes[bg_idx]
        line = slide.shapes[line_idx]
        img = slide.shapes[img_idx]
        t = slide.shapes[t_idx]
        d = slide.shapes[d_idx]

        move(bg, left=x, top=y, width=3.58, height=0.86)
        move(line, left=x + 0.10, top=y + 0.08, width=0.06, height=0.68)
        move(img, left=x + 0.22, top=y + 0.10, width=0.74, height=0.50)
        move(t, left=x + 1.02, top=y + 0.13, width=2.35, height=0.22)
        move(d, left=x + 1.02, top=y + 0.40, width=2.35, height=0.22)
        set_font(t, 15.5, True)
        set_font(d, 14.5, False)


def fix_slide_30(slide):
    # Widen and deepen the answer/constraint panels.
    left_box = slide.shapes[23]
    right_box = slide.shapes[33]
    move(left_box, left=0.72, top=4.18, width=5.98, height=2.12)
    move(right_box, left=6.63, top=4.18, width=5.98, height=2.12)

    for idx in [25, 35]:
        move(slide.shapes[idx], top=4.42, height=0.28)
        set_font(slide.shapes[idx], 17.5, True)
    for idx in [28, 30, 32, 38, 40, 42]:
        set_font(slide.shapes[idx], 15.5, False)

    # Relax the bullet blocks inside the cards.
    left_items = [(27, 28), (29, 30), (31, 32)]
    right_items = [(37, 38), (39, 40), (41, 42)]
    y_positions = [4.90, 5.28, 5.66]
    for (dot_idx, text_idx), yy in zip(left_items, y_positions):
        move(slide.shapes[dot_idx], left=1.26, top=yy, width=0.08, height=0.08)
        move(slide.shapes[text_idx], left=1.44, top=yy - 0.04, width=5.05, height=0.30)
    for (dot_idx, text_idx), yy in zip(right_items, y_positions):
        move(slide.shapes[dot_idx], left=7.16, top=yy, width=0.08, height=0.08)
        move(slide.shapes[text_idx], left=7.34, top=yy - 0.04, width=5.15, height=0.30)

    # Make the top summary cards slightly roomier.
    top_cards = [(4, 5, 6, 7), (8, 9, 10, 11), (12, 13, 14, 15), (16, 17, 18, 19)]
    xs = [0.64, 3.64, 6.64, 9.64]
    widths = [2.95, 2.95, 2.95, 2.95]
    for i, (bg_idx, line_idx, t_idx, d_idx) in enumerate(top_cards):
        x = xs[i]
        move(slide.shapes[bg_idx], left=x, top=2.08, width=2.95, height=1.56)
        move(slide.shapes[line_idx], left=x + 0.10, top=2.15, width=0.08, height=1.40)
        move(slide.shapes[t_idx], left=x + 0.24, top=2.35, width=2.55, height=0.48)
        move(slide.shapes[d_idx], left=x + 0.24, top=2.74, width=2.50, height=1.10)
        set_font(slide.shapes[t_idx], 16.5, True)
        set_font(slide.shapes[d_idx], 14.5, False)


def fix_slide_59(slide):
    # Pull the left rationale card slightly narrower and make the right flow airier.
    move(slide.shapes[5], left=0.55, top=2.02, width=4.02, height=3.10)
    move(slide.shapes[6], left=0.55, top=2.02, width=0.08, height=3.10)
    set_font(slide.shapes[7], 16.5, True)
    move(slide.shapes[7], left=0.70, top=2.10, width=3.45, height=0.30)
    set_font(slide.shapes[8], 16, False)
    set_font(slide.shapes[9], 15.5, False)
    set_font(slide.shapes[10], 16, False)
    set_font(slide.shapes[11], 15.5, False)
    set_font(slide.shapes[12], 16, False)
    set_font(slide.shapes[13], 15.5, False)
    move(slide.shapes[8], left=0.72, top=2.40, width=0.12, height=0.24)
    move(slide.shapes[9], left=0.88, top=2.37, width=3.42, height=0.30)
    move(slide.shapes[10], left=0.72, top=2.68, width=0.12, height=0.24)
    move(slide.shapes[11], left=0.88, top=2.64, width=3.30, height=0.30)
    move(slide.shapes[12], left=0.72, top=2.96, width=0.12, height=0.24)
    move(slide.shapes[13], left=0.88, top=2.92, width=3.42, height=0.30)

    flow = [
        (15, 16, 17, 4.86, 2.38, 1.26, 0.92),
        (19, 20, 21, 6.14, 2.38, 1.26, 0.92),
        (23, 24, 25, 7.42, 2.38, 1.26, 0.92),
        (27, 28, 29, 8.70, 2.38, 1.42, 0.92),
        (31, 32, 33, 10.14, 2.38, 1.24, 0.92),
    ]
    for bg_idx, title_idx, desc_idx, x, y, w, h in flow:
        move(slide.shapes[bg_idx], left=x, top=y, width=w, height=h)
        move(slide.shapes[title_idx], left=x + 0.08, top=y + 0.08, width=w - 0.16, height=0.24)
        move(slide.shapes[desc_idx], left=x + 0.08, top=y + 0.34, width=w - 0.16, height=0.48)
        set_font(slide.shapes[title_idx], 16, True)
        set_font(slide.shapes[desc_idx], 14.5, False)

    move(slide.shapes[35], left=4.84, top=3.48, width=7.28, height=1.15)
    move(slide.shapes[36], left=4.98, top=3.78, width=7.10, height=0.50)
    set_font(slide.shapes[35], 17.5, True)
    set_font(slide.shapes[36], 15.5, False)
    move(slide.shapes[37], left=0.58, top=5.18, width=11.45, height=0.34)
    set_font(slide.shapes[37], 15.5, False)


def fix_slide_60(slide):
    top = [
        (5, 6, 7, 0.78, 2.24, 1.46, 0.96),
        (9, 10, 11, 2.53, 2.24, 1.46, 0.96),
        (13, 14, 15, 4.28, 2.24, 1.46, 0.96),
        (17, 18, 19, 6.03, 2.24, 2.10, 0.96),
        (21, 22, 23, 7.95, 2.24, 1.46, 0.96),
    ]
    for bg_idx, title_idx, desc_idx, x, y, w, h in top:
        move(slide.shapes[bg_idx], left=x, top=y, width=w, height=h)
        move(slide.shapes[title_idx], left=x + 0.10, top=y + 0.08, width=w - 0.20, height=0.22)
        move(slide.shapes[desc_idx], left=x + 0.08, top=y + 0.35, width=w - 0.16, height=0.58)
        set_font(slide.shapes[title_idx], 15.5, True)
        set_font(slide.shapes[desc_idx], 14.5, False)

    # Bottom three cards get more width/height.
    # First lower card
    move(slide.shapes[24], left=0.78, top=4.00, width=2.95, height=1.16)
    move(slide.shapes[25], left=0.78, top=4.00, width=0.08, height=1.16)
    move(slide.shapes[26], left=0.94, top=4.08, width=2.60, height=0.34)
    move(slide.shapes[27], left=1.00, top=4.42, width=0.12, height=0.24)
    move(slide.shapes[28], left=1.15, top=4.38, width=2.55, height=0.30)
    move(slide.shapes[29], left=1.00, top=4.69, width=0.12, height=0.24)
    move(slide.shapes[30], left=1.15, top=4.65, width=2.55, height=0.30)
    set_font(slide.shapes[26], 16.5, True)
    set_font(slide.shapes[28], 15.5, False)
    set_font(slide.shapes[30], 15.5, False)

    # Second lower card
    move(slide.shapes[31], left=3.95, top=4.00, width=3.15, height=1.16)
    move(slide.shapes[32], left=3.95, top=4.00, width=0.08, height=1.16)
    move(slide.shapes[33], left=4.11, top=4.08, width=2.85, height=0.34)
    move(slide.shapes[34], left=4.16, top=4.42, width=0.12, height=0.24)
    move(slide.shapes[35], left=4.31, top=4.38, width=2.82, height=0.30)
    move(slide.shapes[36], left=4.16, top=4.69, width=0.12, height=0.24)
    move(slide.shapes[37], left=4.31, top=4.65, width=2.60, height=0.30)
    set_font(slide.shapes[33], 16.5, True)
    set_font(slide.shapes[35], 15.5, False)
    set_font(slide.shapes[37], 15.5, False)

    # Third lower card
    move(slide.shapes[38], left=7.28, top=4.00, width=3.35, height=1.16)
    move(slide.shapes[39], left=7.28, top=4.00, width=0.08, height=1.16)
    move(slide.shapes[40], left=7.44, top=4.08, width=3.00, height=0.34)
    move(slide.shapes[41], left=7.48, top=4.42, width=0.12, height=0.24)
    move(slide.shapes[42], left=7.63, top=4.38, width=3.05, height=0.30)
    set_font(slide.shapes[40], 16.5, True)
    set_font(slide.shapes[42], 15.5, False)

    move(slide.shapes[43], left=0.68, top=5.47, width=11.95, height=0.38)
    move(slide.shapes[44], left=1.10, top=5.55, width=11.20, height=0.26)
    set_font(slide.shapes[44], 15.5, False)


def fix_slide_61(slide):
    # Top four columns: preserve the original layout but give the cards a little more air.
    moves = [
        (5, 6, 7, 0.85, 2.08, 2.65, 2.52),
        (8, 9, 10, 3.85, 2.08, 2.65, 2.52),
        (11, 12, 13, 6.85, 2.08, 2.65, 2.52),
        (14, 15, 16, 9.85, 2.08, 2.65, 2.52),
    ]
    for bg_idx, accent_idx, head_idx, x, y, w, h in moves:
        move(slide.shapes[bg_idx], left=x, top=y, width=w, height=h)
        move(slide.shapes[accent_idx], left=x + 0.00, top=y, width=0.08, height=h)
        move(slide.shapes[head_idx], left=x + 0.16, top=y + 0.10, width=w - 0.32, height=0.34)
        set_font(slide.shapes[head_idx], 16.5, True)

    pill_groups = [
        [(17, 18), (19, 20), (21, 22)],
        [(23, 24), (25, 26), (27, 28), (29, 30)],
        [(31, 32), (33, 34), (35, 36)],
        [(37, 38), (39, 40), (41, 42)],
    ]
    x_starts = [1.00, 4.00, 7.00, 10.00]
    widths = [2.20, 2.20, 2.20, 2.20]
    y_starts = [2.66, 2.66, 2.66, 2.66]
    for gi, group in enumerate(pill_groups):
        x = x_starts[gi]
        y = y_starts[gi]
        for bg_idx, txt_idx in group:
            move(slide.shapes[bg_idx], left=x, top=y, width=2.25, height=0.40)
            move(slide.shapes[txt_idx], left=x + 0.07, top=y + 0.09, width=widths[gi], height=0.24)
            set_font(slide.shapes[txt_idx], 14.5, False)
            y += 0.50

    # Give the long ingest endpoint a little extra room.
    move(slide.shapes[18], left=1.05, top=2.72, width=2.32, height=0.34)
    set_font(slide.shapes[18], 13.5, False)

    move(slide.shapes[44], left=0.85, top=4.90, width=11.58, height=1.00)
    move(slide.shapes[45], left=1.00, top=5.05, width=11.20, height=0.30)
    set_font(slide.shapes[45], 16.0, True)
    for idx in [47, 49, 51]:
        move(slide.shapes[idx], top=slide.shapes[idx].top.inches + 0.02)
        set_font(slide.shapes[idx], 14.5, False)


def fix_slide_62(slide):
    # Minor cleanup so the already-good summary slide breathes a little more.
    move(slide.shapes[7], left=1.15, top=2.24, width=1.75, height=0.42)
    move(slide.shapes[8], left=7.20, top=2.24, width=2.05, height=0.42)
    for idx in [10, 12, 14, 16, 18, 20, 22, 24]:
        set_font(slide.shapes[idx], 15.0, False)
    for idx in [11, 13, 15, 17, 19, 21, 23]:
        set_font(slide.shapes[idx], 15.0, False)


def main():
    src = Path("external/anduril-sdk介绍_原文仅排版修正版.pptx")
    prs = Presentation(str(src))
    for idx, slide in enumerate(prs.slides, start=1):
        if idx == 5:
            fix_slide_5(slide)
        elif idx == 30:
            fix_slide_30(slide)
        elif idx == 59:
            fix_slide_59(slide)
        elif idx == 60:
            fix_slide_60(slide)
        elif idx == 61:
            fix_slide_61(slide)
        elif idx == 62:
            fix_slide_62(slide)
    prs.save(str(DST_DEFAULT))


if __name__ == "__main__":
    main()
